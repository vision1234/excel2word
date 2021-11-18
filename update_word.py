# -*- coding:utf-8 -*-
"""
@author: 冯伟
@file: update_word.py
@time: 2021/11/16 0:00
"""
import math
from docx import Document
from shutil import copyfile
from pandas import read_excel
from docx.shared import Pt
from datetime import datetime
import os
def read_from_excel(file_name):
    df=read_excel(file_name)
    return df


def convertNumToChinese(totalPrice):
    dictChinese = [u'零',u'壹',u'贰',u'叁',u'肆',u'伍',u'陆',u'柒',u'捌',u'玖']
    unitChinese = [u'',u'拾',u'佰',u'仟','',u'拾',u'佰',u'仟']
    #将整数部分和小数部分区分开
    partA = int(math.floor(totalPrice))
    partB = round(totalPrice-partA, 2)
    strPartA = str(partA)
    strPartB = ''
    if partB != 0:
        strPartB = str(partB)[2:]

    singleNum = []
    if len(strPartA) != 0:
        i = 0
        while i < len(strPartA):
            singleNum.append(strPartA[i])
            i = i+1
    #将整数部分先压再出，因为可以从后向前处理，好判断位数
    tnumChinesePartA = []
    numChinesePartA = []
    j = 0
    bef = '0';
    if len(strPartA) != 0:
        while j < len(strPartA) :
            curr = singleNum.pop()
            if curr == '0' and bef !='0':
                tnumChinesePartA.append(dictChinese[0])
                bef = curr
            if curr != '0':
                tnumChinesePartA.append(unitChinese[j])
                tnumChinesePartA.append(dictChinese[int(curr)])
                bef = curr
            if j == 3:
                tnumChinesePartA.append(u'萬')
                bef = '0'
            j = j+1

        for i in range(len(tnumChinesePartA)):
            numChinesePartA.append(tnumChinesePartA.pop())
    A = ''
    for i in numChinesePartA:
        A = A+i
    #小数部分很简单，只要判断下角是否为零
    B = ''
    if len(strPartB) == 1:
        B = dictChinese[int(strPartB[0])] + u'角'
    if len(strPartB) == 2 and strPartB[0] != '0':
        B = dictChinese[int(strPartB[0])] + u'角' + dictChinese[int(strPartB[1])] + u'分'
    if len(strPartB) == 2 and strPartB[0] == '0':
        B = dictChinese[int(strPartB[0])] + dictChinese[int(strPartB[1])] + u'分'

    if len(strPartB) == 0:
        S = A + u'圆整'
    if len(strPartB)!= 0:
        S = A + u'圆' +B
    return S
def get_date_list(date_str_s,date_str_e):

    res=date_str_s.strftime("%Y/%m/%d").split("/")
    res+=date_str_e.strftime("%Y/%m/%d").split("/")
    for r in res:
        if r.startswith("0"):
            res[res.index(r)]=r[1:]
    return res
def input2word(line,source_name,file_name):
    copyfile(source_name,file_name)
    docnment=Document(file_name)
    docnment.styles['Normal'].font.name = u'微软雅黑'
    docnment.styles['Normal'].font.size =Pt(9)
    table=docnment.tables[0]
    # print(table.cell(7,3).text)
    table.cell(1,0).text="编号：{}".format(line[2])
    date_list=get_date_list(line[6],line[7])
    # print(date_list)
    table.cell(2,0).text="结算单（   {}年 {} 月 {} 日  至  {}  年 {} 月 {} 日 ）".format(*date_list)
    table.cell(5,1).text=line[17]
    table.cell(8,0).text=line[9]
    table.cell(8,1).text=line[5]
    table.cell(8,3).text=str(line[19])
    table.cell(8,4).text=str(line[18])
    table.cell(8,5).text=str(line[20])
    money_uper=convertNumToChinese(line[20])
    table.cell(9,1).text="{}元（大写：{}，含税）".format(str(line[20]),money_uper)
    table.cell(11,1).text=str(line[15])
    docnment.save(file_name)

def jiexi_excel(df,temp_file):
    for v in df .values[2:]:
        input2word(v,"template/{}".format(temp_file),"output/{}.docx".format(v[9]))
# input2word("11月项目结算单.docx", "new_file.docx")
if __name__ == '__main__':
    if not os.path.exists("template"):
        os.mkdir("template")
    if not os.path.exists("input"):
        os.mkdir("input")
    if not os.path.exists("output"):
        os.mkdir("output")
    res0=[f for r,d,f in os.walk("template")]
    res1=[f for r,d,f in os.walk("input")]
    if not len(res0[0]):
        print("template 是空的,放个模板文件进去")
    elif not len(res1[0]):
        print("input 是空的,放个数据源文件进去")
    else:
        print("开搞………………………………………………")
        res=read_from_excel('input/{}'.format(res1[0][0]))
        jiexi_excel(res,res0[0][0])
        print("成了,放进output里了")